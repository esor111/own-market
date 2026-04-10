"""
Backfill replay-safe NRB macro context from discovered current-macro xlsx assets.

Usage:
    python backfill_nrb_macro_context.py 2025
"""
import io
import json
import os
import re
import sys
from datetime import datetime
from glob import glob

import requests
from openpyxl import load_workbook

from config import VALIDATION_DIR


BACKFILL_DIR = os.path.join(VALIDATION_DIR, "historical_context_backfills")
USER_AGENT = "market-gist-macro-backfill/1.0"
TIMEOUT_SECONDS = 60

METRIC_PATTERNS = {
    "remittance_billion": "workers' remittances",
    "interbank_rate_pct": "weighted average interbank rate of commercial banks",
    "private_credit_billion": "bfis credit to private sector",
    "claims_private_sector_yoy_pct": "claims on private sector (y-o-y)",
}


def save_json(path, payload):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", encoding="utf-8") as handle:
        json.dump(payload, handle, indent=2)


def save_text(path, payload):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", encoding="utf-8") as handle:
        handle.write(payload)


def load_json(path):
    with open(path, "r", encoding="utf-8") as handle:
        return json.load(handle)


def _clean_number(value):
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    match = re.search(r"-?\d+(?:\.\d+)?", text.replace(",", ""))
    if not match:
        return None
    return float(match.group(0))


def _load_workbook(url):
    response = requests.get(
        url,
        headers={"User-Agent": USER_AGENT},
        timeout=TIMEOUT_SECONDS,
    )
    response.raise_for_status()
    return load_workbook(io.BytesIO(response.content), read_only=True, data_only=True)


def _extract_metric_row_values(sheet, needle):
    needle = needle.lower()
    for row in sheet.iter_rows(values_only=True):
        values = list(row)
        joined = " | ".join("" if value is None else str(value) for value in values).lower()
        if needle not in joined:
            continue
        numeric_values = [_clean_number(value) for value in values]
        numeric_values = [value for value in numeric_values if value is not None]
        if len(numeric_values) < 2:
            return None
        return {
            "previous_value": numeric_values[-2],
            "current_value": numeric_values[-1],
        }
    return None


def _parse_upload_year_month(url):
    match = re.search(r"/uploads/(\d{4})/(\d{2})/", url)
    if not match:
        return None, None
    return int(match.group(1)), int(match.group(2))


def _shift_month(year, month, delta=-1):
    month_index = (year * 12 + (month - 1)) + delta
    shifted_year = month_index // 12
    shifted_month = (month_index % 12) + 1
    return shifted_year, shifted_month


def _load_discovery_manifests():
    paths = sorted(glob(os.path.join(BACKFILL_DIR, "latest__nrb_current_macro_discovery__*_v1.json")))
    return [load_json(path) for path in paths]


def build_nrb_macro_backfill(target_year):
    manifests = _load_discovery_manifests()
    records = []
    source_fy_tokens = []

    for manifest in manifests:
        fy_token = manifest.get("fy_token")
        source_fy_tokens.append(fy_token)
        for subcategory in manifest.get("subcategories", []):
            tables_post = next((post for post in subcategory.get("posts", []) if post.get("post_type") == "tables" and post.get("asset_type") == "xlsx"), None)
            if not tables_post:
                continue
            upload_year, upload_month = _parse_upload_year_month(tables_post.get("final_url") or "")
            if not upload_year or not upload_month:
                continue
            period_end_year, period_end_month = _shift_month(upload_year, upload_month, delta=-1)
            if period_end_year != target_year:
                continue

            workbook = _load_workbook(tables_post["final_url"])
            if "1.SMIs" not in workbook.sheetnames:
                continue
            sheet = workbook["1.SMIs"]
            metrics = {}
            missing_metrics = []
            for metric_key, needle in METRIC_PATTERNS.items():
                extracted = _extract_metric_row_values(sheet, needle)
                if extracted is None:
                    missing_metrics.append(metric_key)
                else:
                    metrics[metric_key] = extracted

            derived_metrics = {}
            remittance_values = metrics.get("remittance_billion") or {}
            remittance_previous = remittance_values.get("previous_value")
            remittance_current = remittance_values.get("current_value")
            if remittance_previous not in (None, 0) and remittance_current is not None:
                derived_metrics["remittance_yoy_pct"] = round(((remittance_current / remittance_previous) - 1) * 100, 4)

            records.append({
                "fy_token": fy_token,
                "subcategory": subcategory.get("subcategory"),
                "label": subcategory.get("label"),
                "asset_url": tables_post.get("final_url"),
                "asset_content_type": tables_post.get("content_type"),
                "upload_year_month": f"{upload_year:04d}-{upload_month:02d}",
                "period_end_year_month": f"{period_end_year:04d}-{period_end_month:02d}",
                "replay_safe_note": "use upload_year_month as the conservative availability month unless a more exact publication date is captured later",
                "metrics": metrics,
                "derived_metrics": derived_metrics,
                "missing_metrics": missing_metrics,
            })

    records.sort(key=lambda item: (item["period_end_year_month"], item["subcategory"]))

    summary = {
        "schema_version": "1.0",
        "built_at": datetime.now().isoformat(),
        "target_year": target_year,
        "source_type": "nrb_current_macro_tables_xlsx",
        "source_fy_tokens": sorted(set(token for token in source_fy_tokens if token)),
        "record_count": len(records),
        "period_end_year_months": [item["period_end_year_month"] for item in records],
        "records": records,
        "notes": [
            "this is a raw replay-safe macro backfill, not yet a replay-attached context layer",
            "values are extracted from the NRB current macro tables workbook (sheet `1.SMIs`)",
            "upload_year_month is currently the conservative replay-safe availability month",
            "macro regime labels should be derived later, after measurement design is agreed",
        ],
    }
    return summary


def render_markdown(summary):
    lines = [
        "# NRB Macro Context Backfill",
        "",
        f"- built_at: `{summary['built_at']}`",
        f"- target_year: `{summary['target_year']}`",
        f"- source_type: `{summary['source_type']}`",
        f"- source_fy_tokens: `{summary['source_fy_tokens']}`",
        f"- record_count: `{summary['record_count']}`",
        "",
        "## Notes",
        "",
    ]
    for note in summary.get("notes", []):
        lines.append(f"- {note}")

    lines.extend(["", "## Records", ""])
    for item in summary.get("records", []):
        lines.append(f"### {item['period_end_year_month']} — {item['label']}")
        lines.append("")
        lines.append(f"- fy_token: `{item['fy_token']}`")
        lines.append(f"- upload_year_month: `{item['upload_year_month']}`")
        lines.append(f"- asset_url: `{item['asset_url']}`")
        for metric_name, values in (item.get("metrics") or {}).items():
            lines.append(
                f"- `{metric_name}`: previous `{values.get('previous_value')}`, current `{values.get('current_value')}`"
            )
        if item.get("missing_metrics"):
            lines.append(f"- missing_metrics: `{item['missing_metrics']}`")
        lines.append("")
    return "\n".join(lines)


def main():
    if len(sys.argv) < 2:
        print("Usage: python backfill_nrb_macro_context.py YEAR")
        sys.exit(1)

    target_year = int(sys.argv[1])
    summary = build_nrb_macro_backfill(target_year)
    dated_json_path = os.path.join(
        BACKFILL_DIR,
        f"{datetime.now().strftime('%Y-%m-%d')}__{target_year}__nrb_macro_context_backfill_v1.json",
    )
    latest_json_path = os.path.join(BACKFILL_DIR, f"latest__{target_year}__nrb_macro_context_backfill_v1.json")
    dated_md_path = os.path.join(
        BACKFILL_DIR,
        f"{datetime.now().strftime('%Y-%m-%d')}__{target_year}__nrb_macro_context_backfill_v1.md",
    )
    latest_md_path = os.path.join(BACKFILL_DIR, f"latest__{target_year}__nrb_macro_context_backfill_v1.md")
    markdown = render_markdown(summary)
    save_json(dated_json_path, summary)
    save_json(latest_json_path, summary)
    save_text(dated_md_path, markdown)
    save_text(latest_md_path, markdown)
    print(json.dumps({
        "target_year": target_year,
        "record_count": summary["record_count"],
        "latest_json_path": latest_json_path,
        "latest_md_path": latest_md_path,
    }, indent=2))


if __name__ == "__main__":
    main()
